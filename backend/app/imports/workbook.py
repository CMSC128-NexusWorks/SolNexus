"""Read an uploaded .xlsx into plain rows, enforcing the contract's upload limits.

An upload is untrusted input from the network, so the file is checked as a zip archive before openpyxl
parses it: a few kilobytes can decompress to gigabytes (a "zip bomb"), and openpyxl would expand it in
memory. openpyxl itself parses XML through defusedxml (installed in requirements.txt), which refuses
entity-expansion payloads.

Nothing here logs or returns cell values in errors: the workbook holds private client lead data.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time
from typing import Union

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

# Proposed operational limits from the data contract, to validate with SolTech before production.
MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_POPULATED_ROWS = 10_000
# The real workbook is under 1 MiB uncompressed; 100 MiB leaves ample room while bounding memory.
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ZIP_ENTRIES = 1_000

CellValue = Union[str, int, float, bool, datetime, date, time]


class WorkbookError(Exception):
    """The upload cannot be read as a supported workbook. Carries a safe code, never cell contents."""

    def __init__(self, code: str, message: str, status_code: int = 422, details: list[dict] | None = None):
        super().__init__(message)
        self.code = code
        self.message = message
        self.status_code = status_code
        self.details = details or []


@dataclass(frozen=True)
class SheetRow:
    """One worksheet row. cells maps a column letter to its non-empty value, as openpyxl typed it."""

    number: int
    cells: dict[str, CellValue]

    @property
    def is_blank(self) -> bool:
        return not self.cells


@dataclass
class Sheet:
    name: str
    # Every row from 1 to the last populated row, blank rows included so adapters can count them as
    # skipped. Formatted-but-empty rows after the last populated row are not attempted rows and are dropped.
    rows: list[SheetRow] = field(default_factory=list)


def _check_archive(data: bytes) -> None:
    if not zipfile.is_zipfile(io.BytesIO(data)):
        raise WorkbookError("invalid_workbook", "File is not an .xlsx workbook")
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = archive.infolist()
        names = {entry.filename for entry in entries}
        if "xl/workbook.xml" not in names:
            raise WorkbookError("invalid_workbook", "File is not an .xlsx workbook")
        if "xl/vbaProject.bin" in names:
            # .xlsm saved as .xlsx. Macros never run here, but the contract accepts plain .xlsx only.
            raise WorkbookError("invalid_workbook", "Macro-enabled workbooks are not accepted")
        # file_size is the size the archive declares. Python's zipfile stops reading an entry at that
        # size and fails the CRC check if the data is longer, so a lying header cannot exceed it.
        if len(entries) > MAX_ZIP_ENTRIES or sum(entry.file_size for entry in entries) > MAX_UNCOMPRESSED_BYTES:
            raise WorkbookError("workbook_too_large", "Workbook is too large to import", status_code=413)


def _is_populated(value: object) -> bool:
    return value is not None and not (isinstance(value, str) and not value.strip())


def read_workbook(data: bytes) -> list[Sheet]:
    """Parse every worksheet into rows. Raises WorkbookError for an unreadable or oversized workbook."""
    if len(data) > MAX_UPLOAD_BYTES:
        raise WorkbookError("workbook_too_large", "Workbook is larger than 10 MiB", status_code=413)
    _check_archive(data)
    try:
        # read_only streams rows instead of building the whole sheet; data_only returns the value Excel
        # last calculated for a formula cell rather than the formula text.
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises many types (KeyError, BadZipFile, defusedxml errors)
        raise WorkbookError("invalid_workbook", "Workbook could not be read") from exc

    sheets: list[Sheet] = []
    populated_rows = 0
    try:
        for worksheet in workbook.worksheets:
            sheet = Sheet(name=worksheet.title)
            pending_blank: list[SheetRow] = []
            # values_only rows start at column A and row 1, with gaps filled by None, so the tuple index
            # is the column position and the enumeration count is the sheet row number.
            for number, values in enumerate(worksheet.iter_rows(values_only=True), start=1):
                cells = {
                    get_column_letter(index): value
                    for index, value in enumerate(values, start=1)
                    if _is_populated(value)
                }
                row = SheetRow(number=number, cells=cells)
                if row.is_blank:
                    pending_blank.append(row)
                    continue
                populated_rows += 1
                if populated_rows > MAX_POPULATED_ROWS:
                    raise WorkbookError("too_many_rows", f"Workbook has more than {MAX_POPULATED_ROWS} populated rows")
                # Blank rows count only when a populated row follows them.
                sheet.rows.extend(pending_blank)
                pending_blank.clear()
                sheet.rows.append(row)
            sheets.append(sheet)
    except WorkbookError:
        raise
    except Exception as exc:
        raise WorkbookError("invalid_workbook", "Workbook could not be read") from exc
    finally:
        workbook.close()
    return sheets
